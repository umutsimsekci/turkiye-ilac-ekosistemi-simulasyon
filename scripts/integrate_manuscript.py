"""Preserve the September manuscript and build a linked revision and supplement.

Requires python-docx. Source file and output paths are relative to this repository.
The validated calculation release is pinned independently of this editorial commit.
"""
from pathlib import Path
import json, re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R=Path(__file__).resolve().parents[1]
OUT=R/'manuscript';SOURCE=OUT/'source/Makale_23Eylul2026.docx'
URL='https://github.com/umutsimsekci/turkiye-ilac-ekosistemi-simulasyon'
COMMIT='33d88bf72326d19aa1a0c527b9d55f810f0cfcec'
METHOD=[
 'Hesaplar Python 3.12 ile ayrı bütçe, kohort, Ar-Ge ve portföy modüllerinde uygulandı. NumPy rassal sayı üretimi, SciPy geçiş matrisinin üssü, pandas tablo işlemleri, Matplotlib yayın şekilleri, Plotly etkileşimli grafikler ve Streamlit senaryo arayüzü için kullanıldı. Ana analizde 20.000 parametre çekilişi ve 20260922 tohumu korundu. Girdi dosyaları, dağılımlar ve birimler yazılım paketinde açıkça tanımlandı.',
 'Model davranışını görünür kılmak için sabit geçiş matrisi altında 10.000 sanal kişinin 20 yıllık durumları örneklendi. Ayrıca 100 araştırma adayının dört aşamadan geçişi, 10.000 bağımsız tekrarda ardışık binomial çekilişlerle simüle edildi. Her iki süreç için 20261012 tohumu kullanıldı. Bu birinci düzey rastlantısallık analizleri, parametre belirsizliği analizinden ayrıdır; yeni klinik gözlem veya ulusal etki tahmini üretmez. Ek Şekiller S1–S6 model dinamiklerini, gecikmeyi, ortak riski ve karar belirsizliğini gösterir.',
 'Yazılım doğrulamasında makalenin referans sayıları, maliyet özdeşlikleri, durum olasılıklarının korunması, emici ölüm durumu, eşit geçiş hızlarının limiti, proje bütçe koşulları ve arayüzün yeniden hesaplaması test edildi. Etkileşimli ekrandaki deterministik parametre değişiklikleri, makalenin sabit olasılıksal dağılımlarını yeniden kalibre etmez.'
]
ACCESS=f'Yeniden üretilebilir hesaplama kodu, parametre dosyaları, Jupyter defteri ve şekil verileri PHARMACOECOSIM v1.0.0 sürümünde saklandı. Depo: {URL}. Hesaplama sürümü kimliği: {COMMIT}. Depo yazar incelemesi sırasında özeldir; henüz açık erişimli değildir. Veri ve literatür kesimi 22 Eylül 2026, yazılım genişletmesi 12 Ekim 2026 tarihindedir. Matematiksel simülasyon ve yazılım eki, model-kod eşleşmesini ve yeniden çalıştırma adımlarını içerir.'
BLIND='Yeniden üretilebilir hesaplama kodu, parametre dosyaları, Jupyter defteri ve şekil verileri sürümlenmiş özel bir depoda saklandı. Depo bağlantısı ve sürüm kimliği kör değerlendirme amacıyla bu kopyada maskelendi. Kod erişimi için yazar tarafından ayrıca uygun erişim sağlanması gerekir. Veri ve literatür kesimi 22 Eylül 2026, yazılım genişletmesi 12 Ekim 2026 tarihindedir. Matematiksel simülasyon ve yazılım eki, model-kod eşleşmesini ve yeniden çalıştırma adımlarını içerir.'

def scrub(doc):
    doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''

def p(doc,text,style='Normal'):
    return doc.add_paragraph(text,style)

def heading(doc,text):return p(doc,text,'Heading 1')

def table(doc,heads,rows,widths):
    t=doc.add_table(rows=1,cols=len(heads));t.autofit=False
    for cells,values in [(t.rows[0].cells,heads)]+[(t.add_row().cells,row) for row in rows]:
        for i,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Cm(widths[i]);cell.text=str(value)
            for para in cell.paragraphs:
                para.style='TableBody';para.paragraph_format.first_line_indent=Cm(0)
                para.paragraph_format.space_before=Pt(3);para.paragraph_format.space_after=Pt(3)
                for run in para.runs:run.font.size=Pt(10)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6ED');cell._tc.get_or_add_tcPr().append(shade)
                for run in cell.paragraphs[0].runs:run.bold=True
    return t

def mathnodes(items):
    result=[]
    for x in items:
        if isinstance(x,str):
            node=OxmlElement('m:r');t=OxmlElement('m:t');t.text=x;node.append(t)
        elif x[0] in ['sub','sup']:
            node=OxmlElement('m:sSub' if x[0]=='sub' else 'm:sSup')
            for tag,content in [('e',x[1]),(x[0],x[2])]:
                el=OxmlElement('m:'+tag);el.extend(mathnodes(content if isinstance(content,list) else [content]));node.append(el)
        else:
            node=OxmlElement('m:f')
            for tag,content in [('num',x[1]),('den',x[2])]:
                el=OxmlElement('m:'+tag);el.extend(mathnodes(content));node.append(el)
        result.append(node)
    return result

def equation(doc,items):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.first_line_indent=Cm(0)
    math=OxmlElement('m:oMath');math.extend(mathnodes(items));para._p.append(math)

def figure(doc,key,title,note):
    cap=p(doc,title,'Caption');cap.paragraph_format.keep_with_next=True
    para=doc.add_paragraph();para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(3);para.paragraph_format.keep_with_next=True
    shape=para.add_run().add_picture(str(R/'results/figures'/f'{key}.png'),width=Cm(16));shape._inline.docPr.set('descr',title+'. '+note)
    pp=p(doc,note);pp.paragraph_format.first_line_indent=Cm(0)
    for run in pp.runs:run.font.size=Pt(10)

def main():
    OUT.mkdir(exist_ok=True)
    for blind,name in [(False,'Makale_simulasyon_entegre'),(True,'Makale_simulasyon_anonim')]:
        doc=Document(SOURCE)
        anchor=next(x for x in doc.paragraphs if x.text=='III. BULGULAR')
        anchor.insert_paragraph_before('2.6. Yazılım uygulaması ve süreç simülasyonları','Heading 2')
        for text in METHOD:anchor.insert_paragraph_before(text,'Normal')
        ref=next(x for x in doc.paragraphs if x.text=='KAYNAKLAR')
        ref.insert_paragraph_before('Veri ve kod erişimi','Heading 2');ref.insert_paragraph_before(BLIND if blind else ACCESS,'Normal')
        scrub(doc);doc.save(OUT/(name+'.docx'))

    doc=Document(SOURCE)
    for child in list(doc._element.body):
        if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
    scrub(doc);doc.core_properties.title='Matematiksel simülasyon ve yazılım eki'
    p(doc,'Matematiksel simülasyon ve yazılım eki','Title')
    p(doc,'Türkiye ilaç ekosisteminin geliştirilmesi için bütünleşik farmakoekonomik karar modeli','Heading 2')
    p(doc,'Bu ek, ana makalenin karar modellerini çalıştırılabilir simülasyonlara bağlar. Temel sayılar yeniden üretildi; yeni süreç simülasyonları modelin nasıl işlediğini görünür kılar. Gözlem verileri ve varsayımlar kodda ayrı tutulur. Sonuçlar, Türkiye için doğrulanmış etki tahminleri değildir.')
    heading(doc,'1. Modelden koda geçiş')
    table(doc,['Soru','Yöntem','Python modülü'],[
        ['Programın bütçesi','Yıllık nakit akışı ve gecikme','budget.py'],['Tedavinin sağlık değeri','Matris üssü ve kohort ödülleri','cohort.py'],['Araştırma başarısı','Aşama riski ve net bugünkü değer','research.py'],['Proje seçimi','64 bileşimin tam taraması','portfolio.py']],[4.8,7.0,4.2])
    p(doc,'Karar modellerinde sabit 2025 TL kullanılır. Toplam bütçe, Ar-Ge ve portföy çıktıları milyar TL; kohort çıktıları kişi başına TL ve QALY birimindedir. Gösterge CSV dosyalarındaki cari TL ve ABD doları serileri bu tutarlara eklenmez.')
    table(doc,['Hesap','Yeniden üretilen değer'],[
        ['Beş yıllık ek bütçe','−1,680 milyar TL'],['İki yıl gecikmede ek bütçe','+0,620 milyar TL'],['ICER','491.323 TL/QALY'],['Tek aday rNPV','0,837792 milyar TL'],['Temel portföy ve yatırımı','G + B + N1 + P; 7,4 milyar TL']],[9,7])
    p(doc,'Ek Şekiller S1–S6, bu ekin kendi numaralandırmasını kullanır. Ana makaledeki altı şekil ve önceki teknik ekin tablo numaraları korunmuştur.')

    doc.add_page_break();heading(doc,'2. Bütçe etkisinin zamana yayılması')
    p(doc,'Programla ek tedaviye geçen kişi sayısı nₜ; referans fiyat c; fiyat farkı d; izlem gideri m; başlangıç yatırımı F ve yıllık işletme A olarak tanımlanır. Yatırım yıl sıfırda, diğer giderler yıl sonunda yer alır.')
    equation(doc,[('sub','ΔC','t'),' = ',('sub','n','t'),' (−cd + m) + A;     ',('sub','ΔC','0'),' = F'])
    p(doc,'Gecikme, mevcut uygulamaya göre ek kullanım eğrisini öteler. Beş yıllık ufuk sabit olduğu için gecikme, bu ufuk içindeki tasarrufu azaltır. Sabit giderler gecikmeyle ertelenmez. Ek kullanım sıfırsa fiyat için başabaş oranı tanımsızdır.')
    figure(doc,'S2_fiyat_gecikme_yuzeyi','Ek Şekil S2. Fiyat farkı ve gecikmenin bütçe sonucu','Negatif ek maliyet tasarruftur. Gecikmeler tam yıl olarak hesaplanır; siyah çizgi komşu ızgara noktaları arasındaki sıfır maliyet sınırını görsel olarak birleştirir. Kesirli gecikmeler ayrıca simüle edilmedi.')
    p(doc,'Temel senaryoda %25 fiyat farkı ve 160.000 ek hasta-yılı, 4 milyar TL brüt ilaç tasarrufu sağlar. İzlem ve sabit giderlerden sonra net tasarruf 1,68 milyar TL olur. İki yıllık gecikmede aynı program 0,62 milyar TL ek bütçe gerektirir. Bunlar ulusal bütçe hesapları değildir.')

    doc.add_page_break();heading(doc,'3. Kohort ve sanal kişi yolları')
    equation(doc,[('sub','x','k+1'),' = ',('sub','x','k'),'P;     P = ',('sup','e','QΔt')])
    p(doc,'Q, yıllık ilerleme ve ölüm hızlarının üretecidir. Stabil → ilerlemiş, stabil → ölüm ve ilerlemiş → ölüm geçişleri mümkündür. Aday yalnız ilerleme hızını RR ile değiştirir. İlaç maliyeti stabil durumda uygulanır. Maliyet ve QALY için trapez düzeltmesi ve döngü ortası iskonto kullanılır.')
    figure(doc,'S1_kohort_dinamikleri','Ek Şekil S1. Analitik kohort ve 10.000 sanal kişi','Düz eğriler aday, kesikli eğriler standart tedavidir. Sağ panelde açık halkalar sabit parametreler altında tek süreç simülasyonunu gösterir; analitik paylardan en büyük mutlak sapma 1,057 yüzde puandır.')
    figure(doc,'S3_sanal_hasta_yollari','Ek Şekil S3. İlk 30 sanal kişinin yolları','Her satır bir rassal yolu temsil eder. Ölüm emicidir; geri dönüş yoktur. Bunlar hasta verisi değildir. Sanal kişi örneklemesi makalenin analitik ICER hesabını değiştirmez.')

    doc.add_page_break();heading(doc,'4. Parametre belirsizliği ve bilgi değeri')
    equation(doc,['NMB(λ) = λΔQ − ΔC'])
    equation(doc,['EVPI = E[max(0,NMB)] − max(0,E[NMB])'])
    p(doc,'PSA, 20.000 parametre çekilişinden oluşur. Bütçe hesabında fiyat farkı ve kullanım ölçeği Gaussian copula ile ilişkilidir; 0,35 latent normal korelasyonudur. Sağlık modelinde hız ve maliyetler gamma, ilerleme hız oranı lognormal ve fayda ağırlıkları beta dağılımlarıyla temsil edilir. Dağılımlar gözlem verisinden kestirilmedi; örnek varsayımlardır.')
    figure(doc,'S5_karar_belirsizligi','Ek Şekil S5. Maliyet etkililik ve tam bilginin değeri','Sol panelde görsel yoğunluğu azaltmak için her sekizinci çekiliş gösterilir; diğer hesaplarda 20.000 çekilişin tamamı kullanılır. Eşikler Türkiye için resmî kabul edilmiş değerler değildir. EVPI kişi başınadır.')
    table(doc,['Sabit makale PSA sonucu','Değer'],[['Tasarruf olasılığı','%64,27'],['500.000 TL/QALY eşiğinde maliyet etkililik','%50,325'],['Aynı eşikte EVPI','50.470 TL / kişi']],[11,5])
    p(doc,'Arayüzde deterministik girdiler değiştirildiğinde bu sabit PSA sonuçları değişmez. Yeni bir ürün için önce karşılaştırmalı klinik etki, maliyet ve fayda dağılımları belirlenmelidir. Sanal kişi rastlantısallığı ile parametre belirsizliği birbirinin yerine kullanılmaz.')

    doc.add_page_break();heading(doc,'5. Araştırma aşamaları ve ortak risk')
    equation(doc,[('sub','p','ruhsat'),' = 0,65 × 0,35 × 0,60 × 0,90 = 0,12285'])
    p(doc,'Her aşamanın gideri o aşamaya ulaşma olasılığıyla ağırlıklandırılır. Başarı halinde 9–18. yıllarda elde edilen yıllık net nakit akışları %10 reel oranla iskonto edilir. Tek adayın rNPV değeri 0,837792 milyar TL’dir. Gelir başlangıcının gecikmesi maliyet takvimini değiştirmez; 10 yıllık gelir süresini ileri taşır.')
    figure(doc,'S4_arastirma_surec_simulasyonu','Ek Şekil S4. Adayların elenmesi ve ortak başarısızlık riski','Sol: 100 bağımsız aday için 10.000 süreç tekrarı. Kesikli çizgi analitik beklenti 12,285; simülasyon ortalaması 12,2568 onaydır. Sağ: nihai başarı korelasyonu için beta-binomial analitik model; histogramdan ayrı hesaplanmıştır.')
    p(doc,'Bağımsız 10 adayda en az bir başarı olasılığı %73,04’tür. Nihai başarılar arasındaki ortak korelasyon 0,15 olduğunda aynı olasılık %54,23 olur. Aday sayısını artırmak ortak riskleri ortadan kaldırmaz. Korelasyon örnek varsayımdır; Türkiye’deki adaylardan kestirilmedi.')
    p(doc,'Ar-Ge rNPV yatırımcı nakit değeridir; sağlık NMB’siyle aynı ölçüt değildir. Aynı ürünün ticari geliri ve hastaya sağladığı sağlık değeri kontrolsüz biçimde toplanmamalıdır.')

    doc.add_page_break();heading(doc,'6. Bütçeli proje bileşimi')
    equation(doc,['Amaç = Σ ',('sub','x','j'),'[',('sub','p','j'),'(z)(λ',('sub','Q','j'),' − ',('sub','ΔC','j'),') − ',('sub','I','j'),'] − ',('sub','I','P'),'z'])
    p(doc,'Beş proje ve bir platform için 64 ikili seçim değerlendirilir. Yatırım bütçesi aşılmaz; N1 ve N2 için platform zorunludur. Platform R/N1/N2 başarılarını göreli %10 artırır, kendisi QALY üretmez. Model risk nötr beklenen sağlık değerini eniyiler.')
    figure(doc,'S6_portfoy_secimi','Ek Şekil S6. Proje seçim sınırı ve uygun bileşimler','Sol panel, mümkün yatırım tutarlarının tamamında hesaplanan seçim sınırıdır. Sağ panelde temel 8 milyar TL bütçe ve 1 milyon TL/QALY eşiği kullanılır. Gri noktalar bütçe veya platform koşulunu sağlamaz.')
    table(doc,['Kod','Proje','Yatırım milyar TL'],[['G','Kritik eşdeğer','1,0'],['B','Biyobenzer','2,2'],['R','Yeniden konumlandırma','1,3'],['N1','Faz I adayı','3,0'],['N2','Faz II adayı','2,5'],['P','Ortak platform','1,2']],[1.5,10,4.5])
    p(doc,'Temel seçim G + B + N1 + P, yatırım 7,4 ve beklenen net sağlık değeri 3,48513 milyar TL’dir. Platforma ayrı sağlık kazanımı veya Ar-Ge nakit getirisi eklenmez. Adayların korelasyonuna göre riskten kaçınan portföy seçimi bu modelin kapsamı dışındadır.')

    doc.add_page_break();heading(doc,'7. Yeniden üretim ve sürüm kaydı')
    p(doc,ACCESS)
    p(doc,'Python 3.12.14; NumPy 2.5.4, SciPy 1.18.1, pandas 3.0.6, Matplotlib 3.11.2, Plotly 7.1.0 ve Streamlit 1.65.0 ile çalıştırıldı. Tam doğrudan bağımlılıklar requirements-tested.txt içinde, ortam kaydı docs/environment.txt içindedir.')
    table(doc,['İşlem','Depo kökünden komut'],[['Kurulum','python -m pip install -e ".[app,dev]"'],['Tam analiz','python -m pharmacoecosim.run'],['Yerel uygulama','python -m streamlit run app.py'],['Doğrulama','python -m pytest -q']],[4,12])
    p(doc,'config/paper.json girdi kaydı; results/summary.json sayısal özet; results/data şekil tabloları; results/figures PNG/PDF/SVG çıktılarıdır. Çevrimdışı HTML ve GIF tam analizle oluşturulur. Eğitim defterinin sekiz kod hücresi baştan sona hatasız çalıştırıldı. Otuz beş test, referans sayıları ve sınır durumları dahil, yerel ortamda geçti.')
    p(doc,'Bu doğrulama yazılımın iç tutarlılığını gösterir. Türkiye’ye özgü gerçek ürün verileriyle dış doğrulama, bağımsız paydaş kalibrasyonu ve klinik geçerleme yapılmadı. Kullanılan varsayımlar docs/PARAMETRELER.csv, tam denklemler docs/MATEMATIK.md, kaynak künyeleri docs/KAYNAKLAR.md içinde bulunur.')
    p(doc,'Kaynak yöntemler ana makalede yer alan Briggs vd. (2012), Sullivan vd. (2014) ve Husereau vd. (2022) ile birlikte değerlendirilmelidir. Yazılım belgeleri: docs.scipy.org/doc/scipy/reference/generated/scipy.linalg.expm.html ve docs.streamlit.io/develop/api-reference/charts/st.plotly_chart (erişim: 12 Ekim 2026).')
    doc.save(OUT/'Matematiksel_simulasyon_eki.docx')
    (OUT/'KOD_ERISIMI.md').write_text('# Kod erişimi ve sürüm\n\n'+ACCESS+'\n\n'+ '\n\n'.join(METHOD)+'\n',encoding='utf8')
    print('Built 3 DOCX files in manuscript/')

if __name__=='__main__':main()
